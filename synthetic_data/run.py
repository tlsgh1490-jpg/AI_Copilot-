from dataclasses import dataclass
from pathlib import Path

import pandas as pd
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill

from .aggregate import assign_operation_day, build_daily_kpis
from .config import build_config, build_scenarios
from .generator import apply_scenario, build_event_table, generate_hourly_normal
from .qa import run_qa


@dataclass(frozen=True)
class OutputPaths:
    workbook_xlsx: Path
    qa_markdown: Path
    chart_sheet_names: list[str]


def _write_workbook(path: Path, hourly: pd.DataFrame, daily: pd.DataFrame, events: pd.DataFrame, checks: pd.DataFrame, scenarios) -> list[str]:
    with pd.ExcelWriter(path, engine="openpyxl", datetime_format="yyyy-mm-dd hh:mm") as writer:
        hourly.to_excel(writer, sheet_name="시간단위_공정데이터", index=False)
        daily.to_excel(writer, sheet_name="일별_KPI_집계", index=False)
        events.to_excel(writer, sheet_name="이상이벤트_목록", index=False)
        checks.to_excel(writer, sheet_name="QA_상세", index=False)
        chart_sheet_names = _add_chart_sheets(writer, hourly, scenarios)
        for worksheet in writer.book.worksheets:
            worksheet.freeze_panes = "A2"
            worksheet.auto_filter.ref = worksheet.dimensions
            for cell in worksheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1F4E78")
                cell.alignment = Alignment(horizontal="center")
            for column in worksheet.columns:
                letter = column[0].column_letter
                longest = max(len(str(cell.value or "")) for cell in column[:100])
                worksheet.column_dimensions[letter].width = min(max(longest + 2, 12), 28)
    return chart_sheet_names


def _save_qa_markdown(path: Path, checks: pd.DataFrame) -> None:
    lines = ["# Coke Oven Gas 합성데이터 QA 요약", "", "| 점검 항목 | 결과 | 상세 |", "|---|---|---|"]
    for row in checks.itertuples(index=False):
        result = "PASS" if row.passed else "REVIEW"
        lines.append(f"| {row.check_name} | {result} | {row.detail} |")
    lines.extend(["", "- 기준 기간: 2024-01-01 07:00 ~ 2026-01-01 06:00", "- 조업일 집계 기준: 07:00 ~ 익일 07:00", "- 데이터는 실제 원본을 복원하지 않은 가상 합성데이터입니다."])
    path.write_text("\n".join(lines), encoding="utf-8")


def _add_chart_sheets(writer: pd.ExcelWriter, hourly: pd.DataFrame, scenarios) -> list[str]:
    chart_sheet_names: list[str] = []
    seen: set[str] = set()
    for scenario in scenarios:
        if scenario.scenario_id in seen:
            continue
        seen.add(scenario.scenario_id)
        begin = scenario.precursor_start - pd.Timedelta(hours=24)
        end = scenario.recovery_end + pd.Timedelta(hours=24)
        data = hourly.loc[hourly["timestamp"].between(begin, end)]
        driver = scenario.driver_ids[0]
        target = scenario.target_ids[0]
        baseline = data.loc[data["timestamp"] < scenario.precursor_start]
        chart_data = pd.DataFrame({
            "timestamp": data["timestamp"],
            f"{driver}_index": data[driver] / baseline[driver].median() * 100,
            f"{target}_index": data[target] / baseline[target].median() * 100,
            "event_phase": data["event_phase"],
        })
        sheet_name = f"그래프_{scenario.scenario_id.replace('SCN_', '')}"
        chart_data.to_excel(writer, sheet_name=sheet_name, index=False)
        worksheet = writer.book[sheet_name]
        chart = LineChart()
        chart.title = f"{scenario.scenario_id} 대표 이벤트 (기준지수=100)"
        chart.y_axis.title = "relative index"
        chart.x_axis.title = "timestamp"
        values = Reference(worksheet, min_col=2, max_col=3, min_row=1, max_row=len(chart_data) + 1)
        categories = Reference(worksheet, min_col=1, min_row=2, max_row=len(chart_data) + 1)
        chart.add_data(values, titles_from_data=True)
        chart.set_categories(categories)
        chart.height = 12
        chart.width = 26
        worksheet.add_chart(chart, "F2")
        chart_sheet_names.append(sheet_name)
    return chart_sheet_names


def run_generation(output_dir: Path) -> OutputPaths:
    output_dir.mkdir(parents=True, exist_ok=True)
    scenarios = build_scenarios()
    hourly = generate_hourly_normal(build_config())
    for scenario in scenarios:
        hourly = apply_scenario(hourly, scenario)
    hourly = assign_operation_day(hourly)
    daily = build_daily_kpis(hourly)
    events = build_event_table(scenarios)
    checks = run_qa(hourly, daily, events)

    workbook_xlsx = output_dir / "COG_합성데이터_2024_2025.xlsx"
    qa_markdown = output_dir / "COG_생성_QA_요약.md"
    chart_sheet_names = _write_workbook(workbook_xlsx, hourly, daily, events, checks, scenarios)
    _save_qa_markdown(qa_markdown, checks)
    return OutputPaths(workbook_xlsx, qa_markdown, chart_sheet_names)


if __name__ == "__main__":
    paths = run_generation(Path("산출물"))
    print(paths.workbook_xlsx)
